"""Resume processing, structured parsing, and ATS-friendly PDF/DOCX generation."""
from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from backend.live_data import extract_skills, normalize_skills, title_skill

# Regex patterns for contact info extraction
EMAIL_REGEX = re.compile(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+")
PHONE_REGEX = re.compile(r"(?:\+?\d{1,3}[-.\s]?)?(?:\(?\d{3}\)?[-.\s]?)?\d{3}[-.\s]?\d{4}")
URL_REGEX = re.compile(r"https?://(?:www\.)?[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}(?:/\S*)?")
GITHUB_REGEX = re.compile(r"(?:https?://)?(?:www\.)?github\.com/[a-zA-Z0-9_-]+", re.IGNORECASE)
LINKEDIN_REGEX = re.compile(r"(?:https?://)?(?:www\.)?linkedin\.com/in/[a-zA-Z0-9_-]+", re.IGNORECASE)


def extract_text_from_pdf(content: bytes) -> str:
    """Extract plain text from PDF bytes via pypdf."""
    import pypdf
    reader = pypdf.PdfReader(BytesIO(content))
    pages_text = []
    for page in reader.pages:
        txt = page.extract_text()
        if txt:
            pages_text.append(txt)
    return "\n\n".join(pages_text)


def extract_text_from_docx(content: bytes) -> str:
    """Extract plain text from DOCX bytes via python-docx."""
    import docx
    doc = docx.Document(BytesIO(content))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                paragraphs.append(" | ".join(row_text))
    return "\n\n".join(paragraphs)


def extract_text_from_txt(content: bytes) -> str:
    """Extract text from plain text bytes with fallback decoding."""
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        return content.decode("latin-1", errors="ignore")


def extract_resume_text(filename: str, content: bytes) -> str:
    """Dispatch extraction based on file extension with validation."""
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix == "pdf":
        text = extract_text_from_pdf(content)
    elif suffix == "docx":
        text = extract_text_from_docx(content)
    elif suffix == "txt":
        text = extract_text_from_txt(content)
    else:
        raise ValueError("Unsupported format. Please upload a .pdf, .docx, or .txt file.")

    clean_text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if len(clean_text) < 30:
        raise ValueError(
            "The uploaded file does not contain sufficient extractable text. "
            "If this is a scanned/image-only PDF, please upload a machine-readable text document."
        )
    return clean_text


def parse_raw_resume_to_document(text: str) -> dict[str, Any]:
    """Parse unstructured resume text into a normalized ResumeDocument structure."""
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    emails = EMAIL_REGEX.findall(text)
    phones = PHONE_REGEX.findall(text)
    github_links = GITHUB_REGEX.findall(text)
    linkedin_links = LINKEDIN_REGEX.findall(text)
    all_links = URL_REGEX.findall(text)

    # Name is typically the first prominent line
    name = lines[0] if lines and len(lines[0]) < 60 and "@" not in lines[0] else "Candidate"

    detected_skills = extract_skills(text)

    # Segment lines by header keywords
    sections: dict[str, list[str]] = {
        "summary": [],
        "education": [],
        "skills": [],
        "experience": [],
        "projects": [],
        "certifications": [],
        "achievements": [],
    }

    current_sec = "summary"
    header_map = {
        "education": "education",
        "academics": "education",
        "academic background": "education",
        "skills": "skills",
        "technical skills": "skills",
        "technologies": "skills",
        "competencies": "skills",
        "experience": "experience",
        "employment": "experience",
        "work experience": "experience",
        "internships": "experience",
        "projects": "projects",
        "technical projects": "projects",
        "key projects": "projects",
        "certifications": "certifications",
        "certificates": "certifications",
        "achievements": "achievements",
        "honors": "achievements",
        "awards": "achievements",
        "summary": "summary",
        "profile": "summary",
        "about": "summary",
    }

    for line in lines[1:]:
        lower_line = line.lower().strip(":# -")
        if lower_line in header_map:
            current_sec = header_map[lower_line]
            continue
        sections[current_sec].append(line)

    return {
        "personal_info": {
            "name": name,
            "email": emails[0] if emails else "",
            "phone": phones[0] if phones else "",
            "location": "",
            "github": github_links[0] if github_links else "",
            "linkedin": linkedin_links[0] if linkedin_links else "",
            "portfolio": (set(all_links) - set(github_links) - set(linkedin_links)).pop() if (set(all_links) - set(github_links) - set(linkedin_links)) else "",
        },
        "summary": " ".join(sections["summary"][:5]),
        "education": [{"degree": "Degree / Academics", "institution": "\n".join(sections["education"][:4]), "year": "", "score": ""}] if sections["education"] else [],
        "skills": [title_skill(s) for s in detected_skills],
        "experience": [{"role": "Experience", "company": "", "duration": "", "bullets": sections["experience"][:8]}] if sections["experience"] else [],
        "projects": [{"title": "Projects", "tech": "", "bullets": sections["projects"][:8], "link": ""}] if sections["projects"] else [],
        "certifications": sections["certifications"][:6],
        "achievements": sections["achievements"][:6],
        "raw_text": text,
    }


def normalize_resume_input(resume_input: dict[str, Any] | str) -> dict[str, Any]:
    """Ensure resume is in the normalized dictionary format whether passed as dict or raw text."""
    if isinstance(resume_input, str):
        return parse_raw_resume_to_document(resume_input)
    if isinstance(resume_input, dict):
        if "personal_info" in resume_input:
            return resume_input
        if "raw_text" in resume_input:
            return parse_raw_resume_to_document(resume_input["raw_text"])
    return parse_raw_resume_to_document(str(resume_input))


# ============================================================================
# ATS-Friendly PDF Generation (via ReportLab)
# ============================================================================

def generate_resume_pdf(resume_data: dict[str, Any], template_id: str = "classic") -> bytes:
    """
    Generate an ATS-friendly single/multi-page PDF using ReportLab.
    Templates: 'classic' (ATS Classic), 'student' (Student Focus), 'developer' (Developer Focus).
    Follows ATS guidelines: standard fonts, machine-readable layout, no multi-column grids.
    """
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    doc_data = normalize_resume_input(resume_data)
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()

    # Custom ATS-safe styles
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        alignment=1,  # Centered
        textColor=colors.HexColor("#111111")
    )
    contact_style = ParagraphStyle(
        "DocContact",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        alignment=1,
        textColor=colors.HexColor("#333333")
    )
    heading_style = ParagraphStyle(
        "SectionHeading",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=15,
        spaceBefore=8,
        spaceAfter=3,
        textColor=colors.HexColor("#000000")
    )
    body_style = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor("#222222")
    )
    bullet_style = ParagraphStyle(
        "BulletText",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        leftIndent=14,
        textColor=colors.HexColor("#222222")
    )

    elements = []

    # Header / Contact Info
    info = doc_data.get("personal_info", {})
    name = info.get("name") or "Candidate Name"
    elements.append(Paragraph(name.upper(), title_style))
    elements.append(Spacer(1, 4))

    contacts = []
    if info.get("email"):
        contacts.append(info["email"])
    if info.get("phone"):
        contacts.append(info["phone"])
    if info.get("location"):
        contacts.append(info["location"])
    if info.get("linkedin"):
        contacts.append(info["linkedin"].replace("https://", ""))
    if info.get("github"):
        contacts.append(info["github"].replace("https://", ""))
    if info.get("portfolio"):
        contacts.append(info["portfolio"].replace("https://", ""))

    elements.append(Paragraph(" | ".join(contacts) or "Contact Information", contact_style))
    elements.append(Spacer(1, 6))
    elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#999999"), spaceAfter=8))

    def render_summary():
        if doc_data.get("summary"):
            elements.append(Paragraph("PROFESSIONAL SUMMARY", heading_style))
            elements.append(Paragraph(doc_data["summary"], body_style))
            elements.append(Spacer(1, 6))

    def render_education():
        edu_list = doc_data.get("education", [])
        if edu_list:
            elements.append(Paragraph("EDUCATION", heading_style))
            for edu in edu_list:
                if isinstance(edu, dict):
                    deg = edu.get("degree", "")
                    inst = edu.get("institution", "")
                    yr = edu.get("year", "")
                    score = edu.get("score", "")
                    line_parts = [f"<b>{deg}</b>" if deg else "", inst]
                    hdr = " — ".join(p for p in line_parts if p)
                    meta = f" | {score}" if score else ""
                    yr_str = f" ({yr})" if yr else ""
                    elements.append(Paragraph(f"{hdr}{yr_str}{meta}", body_style))
                else:
                    elements.append(Paragraph(str(edu), body_style))
            elements.append(Spacer(1, 6))

    def render_skills():
        skills = doc_data.get("skills", [])
        if skills:
            elements.append(Paragraph("TECHNICAL SKILLS", heading_style))
            if isinstance(skills, dict):
                for cat, sk_list in skills.items():
                    if sk_list:
                        sk_str = ", ".join(sk_list) if isinstance(sk_list, list) else str(sk_list)
                        elements.append(Paragraph(f"<b>{cat.title()}:</b> {sk_str}", body_style))
            elif isinstance(skills, list):
                elements.append(Paragraph(", ".join(str(s) for s in skills), body_style))
            elements.append(Spacer(1, 6))

    def render_experience():
        exp_list = doc_data.get("experience", [])
        if exp_list:
            elements.append(Paragraph("EXPERIENCE & INTERNSHIPS", heading_style))
            for exp in exp_list:
                if isinstance(exp, dict):
                    role = exp.get("role", "")
                    company = exp.get("company", "")
                    dur = exp.get("duration", "")
                    title_line = f"<b>{role}</b>" + (f" — {company}" if company else "") + (f" ({dur})" if dur else "")
                    elements.append(Paragraph(title_line, body_style))
                    for b in exp.get("bullets", []):
                        if b.strip():
                            elements.append(Paragraph(f"• {b.strip()}", bullet_style))
                else:
                    elements.append(Paragraph(str(exp), body_style))
            elements.append(Spacer(1, 6))

    def render_projects():
        proj_list = doc_data.get("projects", [])
        if proj_list:
            elements.append(Paragraph("TECHNICAL PROJECTS", heading_style))
            for p in proj_list:
                if isinstance(p, dict):
                    title = p.get("title", "")
                    tech = p.get("tech", "")
                    link = p.get("link", "")
                    hdr = f"<b>{title}</b>" + (f" | <i>{tech}</i>" if tech else "") + (f" [{link}]" if link else "")
                    elements.append(Paragraph(hdr, body_style))
                    for b in p.get("bullets", []):
                        if b.strip():
                            elements.append(Paragraph(f"• {b.strip()}", bullet_style))
                else:
                    elements.append(Paragraph(str(p), body_style))
            elements.append(Spacer(1, 6))

    def render_certifications():
        certs = doc_data.get("certifications", [])
        if certs:
            elements.append(Paragraph("CERTIFICATIONS", heading_style))
            for c in certs:
                elements.append(Paragraph(f"• {c}", bullet_style))
            elements.append(Spacer(1, 6))

    def render_achievements():
        achievements = doc_data.get("achievements", [])
        if achievements:
            elements.append(Paragraph("ACHIEVEMENTS & LEADERSHIP", heading_style))
            for a in achievements:
                elements.append(Paragraph(f"• {a}", bullet_style))
            elements.append(Spacer(1, 6))

    # Order sections by chosen template
    if template_id == "student":
        # Student template: Education -> Skills -> Projects -> Experience/Internships -> Certifications -> Achievements
        render_education()
        render_skills()
        render_projects()
        render_experience()
        render_certifications()
        render_achievements()
    elif template_id == "developer":
        # Developer template: Summary -> Skills -> Projects -> Experience -> Education -> Certifications
        render_summary()
        render_skills()
        render_projects()
        render_experience()
        render_education()
        render_certifications()
    else:
        # ATS Classic: Summary -> Education -> Skills -> Experience -> Projects -> Certifications -> Achievements
        render_summary()
        render_education()
        render_skills()
        render_experience()
        render_projects()
        render_certifications()
        render_achievements()

    doc.build(elements)
    return buffer.getvalue()


# ============================================================================
# ATS-Friendly DOCX Generation (via python-docx)
# ============================================================================

def generate_resume_docx(resume_data: dict[str, Any], template_id: str = "classic") -> bytes:
    """Generate an ATS-friendly .docx document adhering to clean machine-readable formatting."""
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc_data = normalize_resume_input(resume_data)
    doc = docx.Document()

    # Configure 0.5 inch margins
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    info = doc_data.get("personal_info", {})
    name = info.get("name") or "Candidate Name"

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_name = p_title.add_run(name.upper())
    run_name.font.name = "Arial"
    run_name.font.size = Pt(16)
    run_name.bold = True
    run_name.font.color.rgb = RGBColor(17, 17, 17)

    # Contacts
    contacts = []
    if info.get("email"):
        contacts.append(info["email"])
    if info.get("phone"):
        contacts.append(info["phone"])
    if info.get("location"):
        contacts.append(info["location"])
    if info.get("linkedin"):
        contacts.append(info["linkedin"].replace("https://", ""))
    if info.get("github"):
        contacts.append(info["github"].replace("https://", ""))
    if info.get("portfolio"):
        contacts.append(info["portfolio"].replace("https://", ""))

    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_after = Pt(12)
    run_c = p_contact.add_run(" | ".join(contacts) or "Contact Information")
    run_c.font.name = "Arial"
    run_c.font.size = Pt(9)
    run_c.font.color.rgb = RGBColor(80, 80, 80)

    def add_section_header(title: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(title.upper())
        run.font.name = "Arial"
        run.font.size = Pt(11)
        run.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)

    def docx_summary():
        if doc_data.get("summary"):
            add_section_header("Professional Summary")
            p = doc.add_paragraph(doc_data["summary"])
            p.paragraph_format.space_after = Pt(6)

    def docx_education():
        edu_list = doc_data.get("education", [])
        if edu_list:
            add_section_header("Education")
            for edu in edu_list:
                if isinstance(edu, dict):
                    deg = edu.get("degree", "")
                    inst = edu.get("institution", "")
                    yr = edu.get("year", "")
                    score = edu.get("score", "")
                    p = doc.add_paragraph()
                    r1 = p.add_run(deg + (" — " if inst else ""))
                    r1.bold = True
                    p.add_run(inst)
                    if yr or score:
                        p.add_run(f" ({yr})" if yr else "")
                        p.add_run(f" | {score}" if score else "")
                else:
                    doc.add_paragraph(str(edu))

    def docx_skills():
        skills = doc_data.get("skills", [])
        if skills:
            add_section_header("Technical Skills")
            if isinstance(skills, dict):
                for cat, sk_list in skills.items():
                    if sk_list:
                        p = doc.add_paragraph()
                        r = p.add_run(f"{cat.title()}: ")
                        r.bold = True
                        p.add_run(", ".join(sk_list) if isinstance(sk_list, list) else str(sk_list))
            elif isinstance(skills, list):
                doc.add_paragraph(", ".join(str(s) for s in skills))

    def docx_experience():
        exp_list = doc_data.get("experience", [])
        if exp_list:
            add_section_header("Experience & Internships")
            for exp in exp_list:
                if isinstance(exp, dict):
                    role = exp.get("role", "")
                    company = exp.get("company", "")
                    dur = exp.get("duration", "")
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(2)
                    r = p.add_run(role + (f" — {company}" if company else ""))
                    r.bold = True
                    if dur:
                        p.add_run(f" ({dur})")
                    for b in exp.get("bullets", []):
                        if b.strip():
                            doc.add_paragraph(b.strip(), style="List Bullet")
                else:
                    doc.add_paragraph(str(exp))

    def docx_projects():
        proj_list = doc_data.get("projects", [])
        if proj_list:
            add_section_header("Technical Projects")
            for p_item in proj_list:
                if isinstance(p_item, dict):
                    title = p_item.get("title", "")
                    tech = p_item.get("tech", "")
                    link = p_item.get("link", "")
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(2)
                    r = p.add_run(title)
                    r.bold = True
                    if tech:
                        p.add_run(f" | {tech}")
                    if link:
                        p.add_run(f" [{link}]")
                    for b in p_item.get("bullets", []):
                        if b.strip():
                            doc.add_paragraph(b.strip(), style="List Bullet")
                else:
                    doc.add_paragraph(str(p_item))

    def docx_certifications():
        certs = doc_data.get("certifications", [])
        if certs:
            add_section_header("Certifications")
            for c in certs:
                doc.add_paragraph(str(c), style="List Bullet")

    def docx_achievements():
        achievements = doc_data.get("achievements", [])
        if achievements:
            add_section_header("Achievements & Leadership")
            for a in achievements:
                doc.add_paragraph(str(a), style="List Bullet")

    if template_id == "student":
        docx_education()
        docx_skills()
        docx_projects()
        docx_experience()
        docx_certifications()
        docx_achievements()
    elif template_id == "developer":
        docx_summary()
        docx_skills()
        docx_projects()
        docx_experience()
        docx_education()
        docx_certifications()
    else:
        docx_summary()
        docx_education()
        docx_skills()
        docx_experience()
        docx_projects()
        docx_certifications()
        docx_achievements()

    out_buffer = BytesIO()
    doc.save(out_buffer)
    return out_buffer.getvalue()
